from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

import structlog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from funding.models import FundedCompany

log = structlog.get_logger(__name__)

_COLUMNS = [
    "Company", "Founder / CEO", "Industry", "HQ Location",
    "Delhi NCR Office", "Last Round (Date & Series)", "Amt Raised",
    "Source", "Work Mode", "LinkedIn PM Roles", "LinkedIn Jobs URL",
    "Careers Page",
]

_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_HEADER_FONT = Font(bold=True, size=11)


def _company_to_row(c: FundedCompany) -> list:
    return [
        c.company,
        c.founder_ceo or c.founder_linkedin,
        c.industry,
        c.hq_location,
        c.delhi_ncr_office,
        c.round_display,
        c.amount_raised,
        c.source_url,
        c.work_mode,
        c.linkedin_pm_roles if c.linkedin_pm_roles > 0 else "",
        c.linkedin_jobs_url,
        c.careers_page,
    ]


def export_funding_csv(companies: list[FundedCompany], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = output_dir / f"funded_companies_{ts}.csv"

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(_COLUMNS)
        for c in companies:
            writer.writerow(_company_to_row(c))

    log.info("funding.csv_exported", path=str(path), count=len(companies))
    return path


def export_funding_excel(companies: list[FundedCompany], output_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "funded_companies.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "Recently Funded Companies"

    ws.append(_COLUMNS)
    for cell in ws[1]:
        cell.font = _HEADER_FONT

    for c in companies:
        row = _company_to_row(c)
        ws.append(row)
        # Highlight companies with open PM roles
        if c.linkedin_pm_roles > 0:
            row_idx = ws.max_row
            for cell in ws[row_idx]:
                cell.fill = _GREEN_FILL

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    wb.save(path)
    log.info("funding.excel_exported", path=str(path), count=len(companies))
    return path
