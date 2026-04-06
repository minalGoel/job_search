from __future__ import annotations

from datetime import datetime
from pathlib import Path

import structlog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

log = structlog.get_logger(__name__)

_COLUMNS = [
    "Platform", "Title", "Company", "Location", "Salary", "Posted Date",
    "Skills", "Description", "Apply Link", "Is Duplicate", "Scraped At",
]

_GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_GREY_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
_HEADER_FONT = Font(bold=True)


def _job_to_row(job: dict) -> list:
    skills = job.get("skills", [])
    if isinstance(skills, list):
        skills = ", ".join(skills)
    desc = job.get("description", "")
    if len(desc) > 500:
        desc = desc[:497] + "..."
    return [
        job.get("platform", ""),
        job.get("title", ""),
        job.get("company", ""),
        job.get("location", ""),
        job.get("salary", ""),
        job.get("posted_date", ""),
        skills,
        desc,
        job.get("apply_link", ""),
        "Yes" if job.get("is_duplicate") else "No",
        job.get("scraped_at", ""),
    ]


def export_csv(jobs: list[dict], output_dir: Path) -> Path:
    """Write a CSV file for this run's jobs."""
    import csv

    output_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M")
    path = output_dir / f"jobs_{ts}.csv"

    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(_COLUMNS)
        for job in jobs:
            writer.writerow(_job_to_row(job))

    log.info("export.csv", path=str(path), count=len(jobs))
    return path


def export_excel(all_jobs: list[dict], new_jobs: list[dict], output_dir: Path) -> Path:
    """Write/overwrite the master Excel file with all jobs + highlighted new ones."""
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "jobs_master.xlsx"

    wb = Workbook()

    # --- Sheet 1: All Jobs ---
    ws_all = wb.active
    ws_all.title = "All Jobs"
    ws_all.append(_COLUMNS)
    for cell in ws_all[1]:
        cell.font = _HEADER_FONT

    new_ids = {j.get("id") for j in new_jobs}
    for job in all_jobs:
        row_data = _job_to_row(job)
        ws_all.append(row_data)
        row_idx = ws_all.max_row
        if job.get("is_duplicate"):
            fill = _GREY_FILL
        elif job.get("id") in new_ids:
            fill = _GREEN_FILL
        else:
            fill = None
        if fill:
            for cell in ws_all[row_idx]:
                cell.fill = fill

    # Auto-width (approximate)
    for col in ws_all.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws_all.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

    # --- Sheet 2: New This Run ---
    ws_new = wb.create_sheet("New This Run")
    ws_new.append(_COLUMNS)
    for cell in ws_new[1]:
        cell.font = _HEADER_FONT
    for job in new_jobs:
        ws_new.append(_job_to_row(job))

    wb.save(path)
    log.info("export.excel", path=str(path), total=len(all_jobs), new=len(new_jobs))
    return path
